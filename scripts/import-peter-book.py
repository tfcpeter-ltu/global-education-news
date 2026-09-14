#!/usr/bin/env python3
"""Import Peter's 50-article DOCX collection into the Study Abroad Navigator."""

from __future__ import annotations

import json
import html
import re
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


TITLES = [
    "中國北宋時期，牛津大學已經開始上課了",
    "英國大學只拿 65 分，是不是快被當了？",
    "成績不是最頂尖，就不能去英國讀好學校？",
    "英國 UCAS 2027 申請已開放：現在能填，9 月 1 日才可送出",
    "英國留學怎麼選？排名高，不代表畢業後工作機會一定多",
    "申請英國大學，要考 IELTS for UKVI Academic 嗎？",
    "熱門科系：物理治療與健康專業",
    "英國低齡留學監護人制度全解析",
    "最新留學權威分析：看見 2027 留學趨勢",
    "英國留學政策變向：智庫建議大學名額減少 30%",
    "什麼是英國大學申請 Clearing？",
    "時尚科系：設計、數位工具與永續產業的結合",
    "想讀電影製作？看看英國 Bristol 這個城市",
    "英國學生簽證申請全攻略",
    "英國大學的學徒制 Apprenticeship",
    "英國大學的三明治年 Sandwich Year 是什麼？",
    "英國留學重大變化：畢業後工作簽證縮短至 18 個月",
    "英國大學學費為什麼越來越貴？",
    "SAT 真的回來了：美國名校招生正在重新洗牌",
    "美國大學申請季開跑：Common App 開放後先做這 5 件事",
    "耶魯大學的高教信任危機報告",
    "放射科到底在做什麼？4 個職業一次搞懂",
    "Elon Musk 旗下公司需要哪些科系的人才？",
    "美國留學文章總整理：申請、SAT、AP、2+2 與趨勢",
    "美國大學正在恢復 SAT 審查",
    "什麼是學術英文 Academic English？",
    "另類種族歧視？亞裔學生在美國進不了一流大學",
    "2026 春季美國大學國際學生少了 20%",
    "加拿大學生簽證為什麼突然變難？",
    "加拿大 2027 大學申請：OUAC 9 月開放，先排志願策略",
    "只看世界排名，可能會低估多倫多都會大學 TMU",
    "2027 加拿大名校申請：UBC Okanagan",
    "加拿大實習新規：大專生不再另辦 Co-op 工簽",
    "加拿大知名商學院怎麼選？不是排名越高越適合",
    "2027 是加拿大讀書加移民的好時機嗎？",
    "2026 加拿大最好就業的科系",
    "加拿大留學懶人包：申請、簽證、就業與移民",
    "加拿大學生簽證申請全攻略",
    "澳洲留學：2027 名額凍結後，選校策略正在改變",
    "澳洲留學市場轉向：台灣學生機會上升中",
    "留學實戰精華：澳洲篇",
    "澳洲稅改 2026–27：留學生家庭要重新思考澳洲路線",
    "如何申請新加坡四大公立大學？",
    "SIM 新加坡管理學院與 SMU 新加坡管理大學有什麼不同？",
    "大學的國際性，是為了什麼？",
    "出國留學，先看世界排名還是畢業後留下來的機會？",
    "數學系：未來 AI 世代的重要基礎",
    "中小學交換生：出國前的第一場預演",
    "化學系 vs 化工系：差一個字，人生路線差很多",
    "孩子面對的是新世界，父母不能再用舊答案教育",
]

COUNTRY_KEYS = {
    "英國": "uk",
    "美國": "us",
    "加拿大": "canada",
    "澳洲": "australia",
    "新加坡": "singapore",
    "跨國／教育專題": "global",
}


def compact(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def summary(text: str, limit: int = 108) -> str:
    text = compact(re.sub(r"^[【\[][^】\]]+[】\]]\s*", "", text))
    if len(text) <= limit:
        return text
    cut = text[:limit]
    for mark in "。！？；":
        pos = cut.rfind(mark)
        if pos >= 54:
            return cut[: pos + 1]
    return cut.rstrip("，、：； ") + "…"


def classify(title: str, number: int) -> tuple[str, str, str, list[str]]:
    if any(k in title for k in ("簽證", "移民", "工簽", "留在")):
        return "政策／簽證與就業", "簽證與就業", "留學政策", ["政策", "簽證", "就業"]
    if any(k in title for k in ("申請", "UCAS", "OUAC", "Common App", "Clearing")):
        return "申請攻略／制度", "申請準備", "大學申請", ["申請", "制度", "時程"]
    if any(k in title for k in ("科系", "物理治療", "放射", "數學系", "化學系", "化工系", "時尚", "電影製作")):
        return "科系介紹／職涯", "選科系", "科系與職涯", ["科系", "職涯", "選擇"]
    if number in (45, 47, 48, 49, 50):
        return "Peter 觀點／教育專題", "教育與生涯探索", "教育觀點", ["Peter觀點", "教育", "生涯"]
    if "總整理" in title or "懶人包" in title or "實戰精華" in title:
        return "留學整理／導讀", "留學規劃", "留學規劃", ["懶人包", "留學規劃"]
    return "制度分析／選校", "研究選校", "選校與制度", ["選校", "制度", "留學規劃"]


def paragraph_hyperlink(paragraph, rels) -> str | None:
    for hyperlink in paragraph._p.xpath(".//w:hyperlink"):
        rid = hyperlink.get(qn("r:id"))
        if rid and rid in rels:
            return rels[rid].target_ref
    return None


def paragraph_image(paragraph, rels) -> str | None:
    for blip in paragraph._p.xpath(".//a:blip"):
        rid = blip.get(qn("r:embed"))
        if rid and rid in rels:
            return rels[rid].target_ref
    return None


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: import-peter-book.py INPUT.docx REPOSITORY_ROOT")
    source = Path(sys.argv[1]).resolve()
    root = Path(sys.argv[2]).resolve()
    document = Document(source)
    rels = document.part.rels
    chapters = []

    starts = []
    for idx, paragraph in enumerate(document.paragraphs):
        match = re.fullmatch(r"CHAPTER\s+(\d{2})", compact(paragraph.text))
        if match:
            starts.append((int(match.group(1)), idx))
    if len(starts) != 50:
        raise RuntimeError(f"expected 50 chapters, found {len(starts)}")

    media_dir = root / "public" / "study-abroad" / "images" / "peter-book"
    media_dir.mkdir(parents=True, exist_ok=True)

    with zipfile.ZipFile(source) as archive:
        for offset, (number, start) in enumerate(starts):
            end = starts[offset + 1][1] if offset + 1 < len(starts) else len(document.paragraphs)
            title_idx = next(
                idx for idx in range(start + 1, end)
                if document.paragraphs[idx].style.name == "Heading 1"
            )
            country_label = compact(document.paragraphs[title_idx + 1].text).replace("國家／主題：", "")
            country = COUNTRY_KEYS[country_label]
            original_url = None
            image_target = None
            body = []
            for paragraph in document.paragraphs[title_idx + 1:end]:
                original_url = paragraph_hyperlink(paragraph, rels) or original_url
                image_target = paragraph_image(paragraph, rels) or image_target
                text = compact(paragraph.text)
                if not text or paragraph.style.name in ("Image Caption", "Source Link"):
                    continue
                if text.startswith("國家／主題：") or text.endswith("· 國際教育專文"):
                    continue
                body.append(text)

            if len(body) >= 2:
                a = re.sub(r"\s+", "", body[0])
                b = re.sub(r"\s+", "", body[1])
                if len(a) < len(b) and (a in b or a[: min(45, len(a))] == b[: min(45, len(a))]):
                    body.pop(0)
            if not body or not image_target:
                raise RuntimeError(f"chapter {number:02d} is missing body or image")

            src_name = Path(image_target).name
            suffix = Path(src_name).suffix.lower()
            out_name = f"peter-book-{number:02d}{suffix}"
            with archive.open("word/" + image_target) as source_image, (media_dir / out_name).open("wb") as output_image:
                shutil.copyfileobj(source_image, output_image)

            article_type, stage, major, tags = classify(TITLES[number - 1], number)
            country_tag = "跨國" if country == "global" else country_label
            article = {
                "slug": f"peter-book-{number:02d}",
                "bookNumber": number,
                "bookTitle": "Peter の國際教育講座 第一集 看懂制度，選對世界",
                "country": country,
                "countryLabel": country_label,
                "major": major,
                "title": TITLES[number - 1],
                "kind": "Peter文章",
                "stage": stage,
                "status": "第一集完整收錄",
                "url": f"peter/peter-book-{number:02d}.html",
                "originalUrl": original_url,
                "sourceDate": "2026",
                "articleType": article_type,
                "tags": list(dict.fromkeys([country_tag, *tags])),
                "description": summary(body[0]),
                "image": f"images/peter-book/{out_name}",
                "imageAlt": f"Peter 的國際教育講座第 {number:02d} 篇：{TITLES[number - 1]}",
                "body": body,
            }
            chapters.append(article)

    output = root / "public" / "study-abroad" / "peter-book-data.js"
    encoded = json.dumps(chapters, ensure_ascii=False, separators=(",", ":"))
    output.write_text(
        "(()=>{const rows=" + encoded + ";window.PETER_BOOK_ARTICLES=rows;"
        "const d=window.STUDY_NAV_DATA;if(!d||!Array.isArray(d.articles))return;"
        "rows.forEach(x=>{if(!d.articles.some(a=>a.slug===x.slug))d.articles.push(x);});})();\n",
        encoding="utf-8",
    )
    article_dir = root / "public" / "study-abroad" / "peter"
    article_dir.mkdir(parents=True, exist_ok=True)
    for article in chapters:
        number = article["bookNumber"]
        country_label = "跨國／教育專題" if article["country"] == "global" else article["countryLabel"]
        body_html = "".join(f"<p>{html.escape(paragraph)}</p>" for paragraph in article["body"])
        tags_html = "".join(f"<span>{html.escape(tag)}</span>" for tag in article["tags"])
        if article["originalUrl"]:
            original_html = f'<a class="dark-outline" href="{html.escape(article["originalUrl"])}" target="_blank" rel="noopener">查看 Facebook 原始貼文 ↗</a>'
        else:
            original_html = '<span class="peter-source-pending">原始文件未附 Facebook 直連</span>'
        previous_html = "<span></span>"
        next_html = "<span></span>"
        if number > 1:
            previous = chapters[number - 2]
            previous_html = f'<a href="peter-book-{number - 1:02d}.html"><small>上一篇</small><strong>{html.escape(previous["title"])}</strong></a>'
        if number < len(chapters):
            following = chapters[number]
            next_html = f'<a href="peter-book-{number + 1:02d}.html"><small>下一篇</small><strong>{html.escape(following["title"])}</strong></a>'
        page = f'''<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(article["title"])}｜Peter 的國際教育講座｜留學導航</title><meta name="description" content="{html.escape(article["description"], quote=True)}"><link rel="stylesheet" href="../styles.css"><link rel="stylesheet" href="../mvp.css"></head><body><header class="topbar"><a class="brand" href="../index.html"><span class="brand-mark">SN</span><span><strong>留學導航</strong><small>Study Abroad Navigator</small></span></a><nav class="desktop-nav"><a href="../student-start.html">從這裡開始</a><a href="../universities.html">找大學</a><a href="../majors.html">找科系</a><a href="../apply.html">怎麼申請</a><a href="../peter-library.html">Peter文章</a></nav><a class="nav-cta" href="../peter-library.html">文章庫</a></header><main class="peter-article-page"><article><nav class="peter-breadcrumb" aria-label="麵包屑"><a href="../peter-library.html">Peter 文章庫</a><span>／</span><a href="../peter-library.html?country={article["country"]}">{html.escape(country_label)}</a><span>／</span><span>第 {number:02d} 篇</span></nav><header class="peter-article-head"><span class="overline">PETER'S COLUMN · {html.escape(country_label)}</span><p class="peter-volume">《看懂制度，選對世界》第一集</p><h1>{html.escape(article["title"])}</h1><p class="peter-deck">{html.escape(article["description"])}</p><div class="peter-article-tags">{tags_html}</div></header><figure class="peter-article-figure"><img src="../{html.escape(article["image"])}" alt="{html.escape(article["imageAlt"], quote=True)}"><figcaption>Peter の國際教育講座原始文章配圖</figcaption></figure><aside class="peter-editor-note"><strong>閱讀提醒</strong><p>本文依原始文集收錄，保留 Peter 的觀點與當時資料。涉及政策、招生、簽證、學費及排名的內容，實際申請前仍應查閱當年度官方公告。</p></aside><section class="peter-prose">{body_html}</section><div class="peter-article-actions">{original_html}<a class="primary" href="../peter-library.html?country={article["country"]}">返回 {html.escape(country_label)}文章</a></div><nav class="peter-chapter-nav">{previous_html}{next_html}</nav></article></main><footer><div><strong>留學導航 Study Abroad Navigator</strong><p>Peter 的國際教育講座 × LTU 國際學術中心</p></div><div><p>政策、招生、簽證與學費資訊請以申請當年度官方公告為準。</p></div></footer></body></html>'''
        (article_dir / f"peter-book-{number:02d}.html").write_text(page, encoding="utf-8")

    sitemap_path = root / "public" / "study-abroad" / "sitemap.xml"
    sitemap = sitemap_path.read_text(encoding="utf-8")
    sitemap = re.sub(r"\n?\s*<!-- PETER BOOK START -->.*?<!-- PETER BOOK END -->", "", sitemap, flags=re.S)
    sitemap_entries = "\n".join(
        f"  <url><loc>https://globalednews.com/study-abroad/peter/peter-book-{number:02d}.html</loc><priority>0.7</priority></url>"
        for number in range(1, 51)
    )
    sitemap = sitemap.replace("</urlset>", f"  <!-- PETER BOOK START -->\n{sitemap_entries}\n  <!-- PETER BOOK END -->\n</urlset>")
    sitemap_path.write_text(sitemap, encoding="utf-8")
    print(f"generated {len(chapters)} articles at {output}")


if __name__ == "__main__":
    main()
